# S2 (review 2026-09-23): objects cut by the image/stack border (or by alignment
# padding, or by the ROI) must be flagged in every mode and kept out of primary.
import numpy as np
import pytest
import tifffile

from Batch import process_condensates
from postprocessing import _prepare_reporting_frames

SHAPE = (8, 40, 40)


def _stack(boxes):
    raw = np.full(SHAPE, 100, dtype=np.uint16)
    mask = np.zeros(SHAPE, dtype=np.uint8)
    for z, y, x in boxes:
        mask[z, y, x] = 1
        raw[z, y, x] = 800
    return raw, mask


def _run(folder, raw, mask, roi=None, name="img", **kwargs):
    folder.mkdir(parents=True, exist_ok=True)
    tifffile.imwrite(folder / f"{name}.tif", raw)
    tifffile.imwrite(folder / f"{name}_Mask.tif", mask)
    return process_condensates(
        tif_path=folder / f"{name}.tif", mask_path=folder / f"{name}_Mask.tif", mode="3d",
        min_voxels=1, auto_roi=roi is None,
        request_roi_func=(lambda shape, is_3d: roi) if roi is not None else None,
        pixel_size_nm=100.0, z_step_nm=300.0, signal_channel=0, dapi_channel=0, **kwargs,
    )


INTERIOR = (slice(3, 5), slice(18, 22), slice(18, 22))
XY_EDGE = (slice(3, 5), slice(0, 4), slice(18, 22))       # y starts at 0
Z_EDGE = (slice(0, 2), slice(30, 34), slice(5, 9))         # z starts at 0


def test_edge_flag_without_mode_a_and_excluded_from_primary(tmp_path):
    raw, mask = _stack([INTERIOR, XY_EDGE, Z_EDGE])
    df = _run(tmp_path, raw, mask)                       # Mode A disabled
    by_y = {round(r.Y_px): bool(r.touches_image_edge) for r in df.itertuples()}
    assert by_y == {20: False, 2: True, 32: True}

    _, primary, excluded, _, _ = _prepare_reporting_frames(df, 0.0, 1e6)
    assert sorted(round(v) for v in primary["Y_px"]) == [20]
    assert len(excluded) == 2


def test_padding_border_counts_as_edge(tmp_path):
    """After expand_canvas the original FOV border is inside the array (bbox x = pad)."""
    raw, mask = _stack([INTERIOR, (slice(3, 5), slice(10, 14), slice(0, 4))])  # 2nd touches x = 0
    pad = ((0, 0), (0, 0), (3, 0))
    folder = tmp_path / "aligned"
    folder.mkdir()
    tifffile.imwrite(folder / "img_alignment_valid.tif", np.pad(np.ones(SHAPE, np.uint8), pad))
    df = _run(folder, np.pad(raw, pad), np.pad(mask, pad), mode_a_enabled=False)
    by_y = {round(r.Y_px): bool(r.touches_image_edge) for r in df.itertuples()}
    assert by_y == {20: False, 12: True}


def test_padding_without_valid_mask_is_not_detected(tmp_path):
    """Documents the limitation for stacks aligned by an older ExQt (no valid-voxel mask)."""
    raw, mask = _stack([(slice(3, 5), slice(10, 14), slice(0, 4))])
    pad = ((0, 0), (0, 0), (3, 0))
    df = _run(tmp_path, np.pad(raw, pad), np.pad(mask, pad))
    assert not df["touches_image_edge"].iloc[0]


def test_object_cut_by_roi_is_flagged_and_excluded(tmp_path):
    raw, mask = _stack([INTERIOR, (slice(3, 5), slice(5, 9), slice(5, 9))])
    roi = np.zeros(SHAPE, dtype=np.int64)
    roi[4, 3:30, 3:30] = 1
    roi[4, 3:30, 3:7] = 0                              # ROI boundary cuts the 2nd object at x = 7
    df = _run(tmp_path, raw, mask, roi=roi)
    cut = df[df["Y_px"] < 10]
    inside = df[df["Y_px"] > 15]
    assert cut["touches_roi_edge"].tolist() == [True]
    assert inside["touches_roi_edge"].tolist() == [False]
    _, primary, _, _, _ = _prepare_reporting_frames(df, 0.0, 1e6)
    assert sorted(round(v) for v in primary["Y_px"]) == [20]


def test_mode_a_columns_match_general_flags(tmp_path):
    raw, mask = _stack([(slice(2, 6), slice(12, 28), slice(12, 28)), (slice(2, 6), slice(0, 10), slice(12, 28))])
    df = _run(tmp_path, raw, mask, mode_a_enabled=True, mode_a_min_core_voxels=5)
    assert (df["mode_a_object_touches_edge"] == df["touches_image_edge"]).all()
    assert (df["mode_a_object_touches_roi_edge"] == df["touches_roi_edge"]).all()
    assert df["touches_image_edge"].any()


def test_excel_summary_counts_boolean_edge_flags(tmp_path):
    from openpyxl import load_workbook
    from postprocessing import generate_excel_stats

    raw, mask = _stack([INTERIOR, XY_EDGE, Z_EDGE])
    df = _run(tmp_path, raw, mask)
    csv = tmp_path / "run_Output_Batch_3d.csv"
    df.to_csv(csv, index=False)
    result = generate_excel_stats(str(csv), min_size=0.0, max_size=1e6)
    ws = load_workbook(result["excel"])["Summary"]
    formulas = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    assert formulas["Touches image edge"].endswith(",TRUE)")
