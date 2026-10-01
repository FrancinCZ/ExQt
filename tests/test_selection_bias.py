# S8 (review 2026-09-23): the primary set needs a valid core FA, which in practice keeps
# only large objects; the report must state this size selection with numbers.
import numpy as np
import pandas as pd

from postprocessing import _prepare_reporting_frames, primary_selection_summary


def _frame():
    big = dict(volume_px=800, volume_bio_um3=0.8, A_object=0.3, A_shell=0.3, A_middle=0.3, A_core=0.3,
               A_object_valid=True, A_shell_valid=True, A_middle_valid=True, A_core_valid=True,
               mode_a_primary_include=True, filename="a.tif", cell_id=1)
    small = dict(big, volume_px=60, volume_bio_um3=0.06, A_core=np.nan, A_core_valid=False,
                 mode_a_primary_include=False)
    frame = pd.DataFrame([big, big, dict(big, volume_px=1200, volume_bio_um3=1.2), small, small, small])
    return frame.assign(object_id=range(1, len(frame) + 1), mode="3d")


def test_primary_selection_summary_reports_size_shift():
    _, _, _, report_df, size_col = _prepare_reporting_frames(_frame(), 0.0, 10.0)
    summary = primary_selection_summary(report_df, size_col)
    assert summary["n_primary"] == 3 and summary["n_excluded_eligible"] == 3
    assert summary["median_px_primary"] == 800 and summary["median_px_excluded"] == 60
    assert summary["primary_fraction"] == 0.5
    assert "larger objects" in summary["message"]


def test_excel_summary_contains_selection_bias_row(tmp_path):
    from openpyxl import load_workbook
    from postprocessing import generate_excel_stats

    csv = tmp_path / "run_Output_Batch_3d.csv"
    _frame().to_csv(csv, index=False)
    result = generate_excel_stats(str(csv), min_size=0.0, max_size=10.0)
    ws = load_workbook(result["excel"])["Summary"]
    rows = {ws.cell(row=r, column=5).value: ws.cell(row=r, column=6).value for r in range(1, ws.max_row + 1)}
    assert "800" in rows["Primary selection bias"] and "60" in rows["Primary selection bias"]
    assert rows["Red flags"] is not None  # rows above were not overwritten by the per-cell table
