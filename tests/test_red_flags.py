# S5 (review 2026-09-23, rules.md "Okamžité hlášení anomálií"): > 1000 objects per
# nucleus or a median object size < 50 voxels must be reported loudly (log, Excel, metadata).
import json
import os

import numpy as np
import pandas as pd
import tifffile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from postprocessing import (
    RED_FLAG_MAX_OBJECTS_PER_NUCLEUS,
    RED_FLAG_MIN_MEDIAN_VOXELS,
    detect_red_flags,
)


def test_thresholds_follow_rules_md():
    assert RED_FLAG_MAX_OBJECTS_PER_NUCLEUS == 1000
    assert RED_FLAG_MIN_MEDIAN_VOXELS == 50


def test_too_many_objects_in_one_nucleus_is_flagged():
    df = pd.DataFrame({
        "filename": ["a.tif"] * 1001 + ["a.tif"] * 1000,
        "cell_id": [1] * 1001 + [2] * 1000,
        "volume_px": [200] * 2001,
    })
    flags = detect_red_flags(df)
    assert [(f["flag"], f["cell_id"], f["value"]) for f in flags] == [("objects_per_nucleus", 1, 1001)]


def test_small_median_size_is_flagged_per_file():
    df = pd.DataFrame({
        "filename": ["small.tif"] * 3 + ["ok.tif"] * 3,
        "cell_id": [1] * 6,
        "volume_px": [8, 12, 60, 80, 90, 100],
    })
    flags = detect_red_flags(df)
    assert [(f["flag"], f["filename"], f["value"]) for f in flags] == [("median_size_below_min", "small.tif", 12.0)]
    assert "min_voxels" in flags[0]["message"]


def test_2d_uses_pixels_and_clean_data_has_no_flags():
    df = pd.DataFrame({"filename": ["a.tif"] * 3, "cell_id": [1] * 3, "area_px": [10, 20, 30]})
    assert detect_red_flags(df)[0]["unit"] == "pixels"
    assert detect_red_flags(df.assign(area_px=[100, 200, 300])) == []


def test_worker_reports_red_flags_in_log_and_metadata(tmp_path):
    from App import AnalysisWorker

    raw = np.full((6, 30, 30), 100, dtype=np.uint16)
    mask = np.zeros((6, 30, 30), dtype=np.uint8)
    for y in (5, 12, 19):
        mask[2, y:y + 2, 10:12] = 1                    # 4-voxel specks: typical over-segmentation
        raw[2, y:y + 2, 10:12] = 700
    tifffile.imwrite(tmp_path / "img.tif", raw)
    tifffile.imwrite(tmp_path / "img_Mask.tif", mask)
    worker = AnalysisWorker({
        "input_folder": str(tmp_path), "output_folder": str(tmp_path / "out"), "mode": "3d",
        "expansion_factor": 1.0, "min_voxels": 1, "auto_roi": True, "review_each_image": False,
        "show_napari": False, "generate_reports": True, "report_excel": True,
        "report_primary_csv": False, "report_excluded_csv": False, "report_raw_audit_csv": False,
        "report_standard_plots": False, "report_mode_a_plots": False, "report_partitioning_plots": False,
        "plot_min_size": 0.0, "plot_max_size": 1e6,
        "pixel_size_nm": 100.0, "z_step_nm": 300.0, "calibration_source": "gui",
        "signal_channel": 0, "dapi_channel": 0,
    })
    messages = []
    worker.progress.connect(messages.append)
    worker.run()
    out = tmp_path / "out"
    meta = json.loads(next(out.glob("*_metadata.json")).read_text(encoding="utf-8"))
    assert [f["flag"] for f in meta["red_flags"]] == ["median_size_below_min"]
    assert any(m.startswith("RED FLAG") for m in messages)

    from openpyxl import load_workbook
    ws = load_workbook(next(out.glob("*_Detailed_Stats.xlsx")))["Summary"]
    values = {ws.cell(row=r, column=5).value: ws.cell(row=r, column=6).value for r in range(1, ws.max_row + 1)}
    assert "median_size_below_min" in values["Red flags"]
